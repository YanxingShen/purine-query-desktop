# -*- coding: utf-8 -*-
"""从 purine-db.json 导出 purine-db.xlsx"""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")

with open(os.path.join(DATA, "purine-db.json"), encoding="utf-8") as f:
    db = json.load(f)

wb = Workbook()
ws = wb.active
ws.title = "嘌呤食物数据库"

headers = ["id","名称","别名","大类","小类","嘌呤(mg/100g)","范围下限","范围上限","等级","食用状态","来源","备注"]
ws.append(headers)
for c in ws[1]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="DDEBF7")

LEVEL_FILL = {
    "低": "C6EFCE", "中": "FFEB9C", "高": "FFC7CE", "极高": "FF6B6B", "未知": "D9D9D9"
}

for it in db["items"]:
    rng = it.get("purine_range") or [None, None]
    aliases = "、".join(it.get("aliases") or [])
    sources = " ; ".join(it.get("sources") or [])
    ws.append([
        it["id"], it["name"], aliases, it["category"], it.get("subcategory",""),
        it["purine_mg_per_100g"] if it["purine_mg_per_100g"] is not None else "暂无数据",
        rng[0] if rng[0] is not None else "",
        rng[1] if rng[1] is not None else "",
        it["level"], it.get("preparation",""), sources, it.get("note","")
    ])
    # 等级着色
    fill = PatternFill("solid", fgColor=LEVEL_FILL.get(it["level"], "FFFFFF"))
    ws.cell(row=ws.max_row, column=9).fill = fill

# 列宽
widths = [8,18,20,12,12,16,10,10,8,14,50,40]
for i,w in enumerate(widths,1):
    ws.column_dimensions[chr(64+i)].width = w
ws.freeze_panes = "A2"

out = os.path.join(DATA, "purine-db.xlsx")
wb.save(out)
print(f"Exported {len(db['items'])} rows to {out}")
